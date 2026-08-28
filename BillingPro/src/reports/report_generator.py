"""
BillingPro - Report Generator
=============================
Generate PDF and Excel reports.
"""

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from datetime import datetime
from typing import List, Dict


class ReportGenerator:
    """Generate printable reports in PDF and Excel formats."""

    @staticmethod
    def generate_pdf(data, output_path: str, title: str = "Report"):
        """Generate PDF report."""
        doc = SimpleDocTemplate(
            output_path,
            pagesize=A4,
            rightMargin=15*mm,
            leftMargin=15*mm,
            topMargin=15*mm,
            bottomMargin=15*mm
        )

        styles = getSampleStyleSheet()
        elements = []

        # Title
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=20,
            textColor=colors.HexColor('#2C3E50'),
            spaceAfter=20,
            alignment=TA_CENTER
        )
        elements.append(Paragraph(title, title_style))

        # Date
        date_style = ParagraphStyle(
            'DateStyle',
            parent=styles['Normal'],
            fontSize=10,
            textColor=colors.grey,
            alignment=TA_CENTER,
            spaceAfter=20
        )
        elements.append(Paragraph(f"Generated on: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}", date_style))

        # Determine data structure and create table
        if isinstance(data, dict):
            # Handle nested data (like daily report with summary and bills)
            if 'summary' in data:
                summary = data.get('summary', {})
                summary_data = [["Metric", "Value"]]
                for key, value in summary.items():
                    if isinstance(value, (int, float)):
                        summary_data.append([key.replace('_', ' ').title(), f"₹{value:,.2f}" if 'amount' in key or 'sales' in key or 'collection' in key else str(value)])

                summary_table = Table(summary_data, colWidths=[80*mm, 80*mm])
                summary_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2C3E50')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 1), (-1, -1), 10),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                    ('TOPPADDING', (0, 0), (-1, -1), 8),
                ]))
                elements.append(summary_table)
                elements.append(Spacer(1, 20))

            # Bills/Details table
            if 'bills' in data:
                bills = data['bills']
                if bills:
                    elements.append(Paragraph("Details", styles['Heading2']))
                    elements.append(Spacer(1, 10))

                    bill_headers = list(bills[0].keys())[:6]  # First 6 columns
                    bill_data = [bill_headers]
                    for bill in bills:
                        row = [str(bill.get(h, '')) for h in bill_headers]
                        bill_data.append(row)

                    col_width = 160*mm / len(bill_headers)
                    bill_table = Table(bill_data, colWidths=[col_width]*len(bill_headers))
                    bill_table.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498DB')),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('FONTSIZE', (0, 0), (-1, 0), 10),
                        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                        ('FONTSIZE', (0, 1), (-1, -1), 9),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                        ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ]))
                    elements.append(bill_table)

            elif 'daily_breakdown' in data:
                daily = data['daily_breakdown']
                if daily:
                    elements.append(Paragraph("Daily Breakdown", styles['Heading2']))
                    elements.append(Spacer(1, 10))

                    headers = list(daily[0].keys())
                    table_data = [headers]
                    for day in daily:
                        row = [str(day.get(h, '')) for h in headers]
                        table_data.append(row)

                    col_width = 160*mm / len(headers)
                    table = Table(table_data, colWidths=[col_width]*len(headers))
                    table.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498DB')),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('FONTSIZE', (0, 0), (-1, 0), 10),
                        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                        ('FONTSIZE', (0, 1), (-1, -1), 9),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                        ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ]))
                    elements.append(table)

        elif isinstance(data, list):
            # Simple list of dictionaries
            if data:
                headers = list(data[0].keys())
                table_data = [headers]
                for item in data:
                    row = [str(item.get(h, '')) for h in headers]
                    table_data.append(row)

                col_width = 160*mm / len(headers)
                table = Table(table_data, colWidths=[col_width]*len(headers))
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2C3E50')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 10),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 1), (-1, -1), 9),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                    ('TOPPADDING', (0, 0), (-1, -1), 6),
                ]))
                elements.append(table)

        # Footer
        elements.append(Spacer(1, 30))
        footer_style = ParagraphStyle(
            'Footer',
            parent=styles['Normal'],
            fontSize=8,
            textColor=colors.grey,
            alignment=TA_CENTER
        )
        elements.append(Paragraph("Generated by BillingPro - Billing Software for Retail Shops", footer_style))

        doc.build(elements)
        return output_path

    @staticmethod
    def generate_excel(data, output_path: str, title: str = "Report"):
        """Generate Excel report."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"

        # Title
        ws['A1'] = title
        ws['A1'].font = Font(size=16, bold=True, color="2C3E50")
        ws.merge_cells('A1:F1')
        ws['A1'].alignment = Alignment(horizontal='center')

        ws['A2'] = f"Generated on: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}"
        ws['A2'].font = Font(size=10, color="7F8C8D")
        ws.merge_cells('A2:F2')
        ws['A2'].alignment = Alignment(horizontal='center')

        # Data
        if isinstance(data, dict):
            if 'bills' in data:
                bills = data['bills']
                if bills:
                    row = 4
                    headers = list(bills[0].keys())[:6]
                    for col, header in enumerate(headers, 1):
                        cell = ws.cell(row=row, column=col, value=header.replace('_', ' ').title())
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
                        cell.alignment = Alignment(horizontal='center')

                    for bill in bills:
                        row += 1
                        for col, header in enumerate(headers, 1):
                            value = bill.get(header, '')
                            cell = ws.cell(row=row, column=col, value=value)
                            cell.alignment = Alignment(horizontal='left' if col == 1 else 'center')

            elif 'daily_breakdown' in data:
                daily = data['daily_breakdown']
                if daily:
                    row = 4
                    headers = list(daily[0].keys())
                    for col, header in enumerate(headers, 1):
                        cell = ws.cell(row=row, column=col, value=header.replace('_', ' ').title())
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
                        cell.alignment = Alignment(horizontal='center')

                    for day in daily:
                        row += 1
                        for col, header in enumerate(headers, 1):
                            value = day.get(header, '')
                            cell = ws.cell(row=row, column=col, value=value)

        elif isinstance(data, list) and data:
            row = 4
            headers = list(data[0].keys())
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header.replace('_', ' ').title())
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
                cell.alignment = Alignment(horizontal='center')

            for item in data:
                row += 1
                for col, header in enumerate(headers, 1):
                    value = item.get(header, '')
                    cell = ws.cell(row=row, column=col, value=value)

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(output_path)
        return output_path
